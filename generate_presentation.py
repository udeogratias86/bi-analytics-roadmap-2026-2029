from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("Roadmap_BI_Analytics_2026_2029.pptx")
W, H = Inches(13.333), Inches(7.5)
BG = "071F1D"
PANEL = "103632"
GREEN = "0E5B4F"
TEAL = "00A99D"
BLUE = "2F80ED"
VIOLET = "7856D9"
ORANGE = "F2994A"
WHITE = "F7FAFC"
MUTED = "B6CBC7"


def color(value):
    return RGBColor.from_string(value)


def shape(slide, kind, x, y, w, h, fill, radius=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color(fill)
    s.line.fill.background()
    return s


def text(slide, value, x, y, w, h, size=18, bold=False, fill=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.text = value
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color(fill)
    return box


def base(prs, title, number=None, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color(BG)
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.14), TEAL)
    if number:
        text(slide, number, Inches(0.55), Inches(0.35), Inches(0.8), Inches(0.3), 10, True, TEAL)
    text(slide, title, Inches(0.55), Inches(0.68), Inches(12), Inches(0.55), 28, True)
    if subtitle:
        text(slide, subtitle, Inches(0.55), Inches(1.25), Inches(11.8), Inches(0.35), 12, False, MUTED)
    text(slide, "BI & ANALYTICS  •  2026—2029", Inches(9.8), Inches(7.05), Inches(2.9), Inches(0.2), 8, True, MUTED, PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, accent=TEAL, icon=None):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, PANEL, True)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(0.07), h, accent)
    if icon:
        shape(slide, MSO_SHAPE.OVAL, x + Inches(0.25), y + Inches(0.25), Inches(0.46), Inches(0.46), accent)
        text(slide, icon, x + Inches(0.25), y + Inches(0.26), Inches(0.46), Inches(0.4), 15, True, WHITE, PP_ALIGN.CENTER)
        tx = x + Inches(0.87)
        tw = w - Inches(1.08)
    else:
        tx, tw = x + Inches(0.28), w - Inches(0.5)
    text(slide, heading, tx, y + Inches(0.2), tw, Inches(0.35), 14, True)
    text(slide, body, tx, y + Inches(0.65), tw, h - Inches(0.85), 10.5, False, MUTED)


def bullets(items):
    return "\n".join("• " + item for item in items)


def make():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H, BG)
    shape(slide, MSO_SHAPE.OVAL, Inches(8.8), Inches(-1.4), Inches(5.6), Inches(5.6), GREEN)
    shape(slide, MSO_SHAPE.OVAL, Inches(9.7), Inches(0.1), Inches(3.8), Inches(3.8), VIOLET)
    shape(slide, MSO_SHAPE.OVAL, Inches(10.5), Inches(1.0), Inches(2.0), Inches(2.0), TEAL)
    text(slide, "ROADMAP TECHNOLOGIQUE", Inches(0.7), Inches(1.35), Inches(5.6), Inches(0.3), 13, True, TEAL)
    text(slide, "BI & Analytics\n2026—2029", Inches(0.7), Inches(1.75), Inches(7.3), Inches(1.8), 42, True)
    text(slide, "Construire une BI industrialisée, centrée produit,\nau service de la performance.", Inches(0.75), Inches(3.85), Inches(6.7), Inches(0.65), 17, False, MUTED)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(5.15), Inches(3.05), Inches(0.6), TEAL, True)
    text(slide, "VISION • VALEUR • EXCELLENCE", Inches(0.88), Inches(5.26), Inches(2.8), Inches(0.25), 10, True, WHITE, PP_ALIGN.CENTER)

    slide = base(prs, "Pourquoi changer ?", "01", "Le constat aujourd’hui : une BI qui doit redevenir un levier de confiance.")
    card(slide, Inches(.65), Inches(1.95), Inches(3.8), Inches(3.55), "Rapports isolés", bullets(["Multiplication des versions", "Vision fragmentée des indicateurs", "Faible réutilisation"]), BLUE, "↗")
    card(slide, Inches(4.78), Inches(1.95), Inches(3.8), Inches(3.55), "Manque de confiance", bullets(["Règles métier peu explicites", "Qualité inégale", "Traçabilité insuffisante"]), ORANGE, "!")
    card(slide, Inches(8.91), Inches(1.95), Inches(3.8), Inches(3.55), "Coûts élevés", bullets(["Développements redondants", "Maintenance complexe", "Time-to-insight trop long"]), VIOLET, "€")
    text(slide, "IMPACT BUSINESS  —  Décisions ralenties  •  Risques accrus  •  Valeur data sous-exploitée", Inches(.75), Inches(6.05), Inches(11.8), Inches(.45), 15, True, TEAL, PP_ALIGN.CENTER)

    slide = base(prs, "Quelle cible ?", "02", "Vision 2029 : une plateforme de décisions fiable, accessible et à l’échelle.")
    card(slide, Inches(.7), Inches(2.0), Inches(3.75), Inches(3.45), "Gold Data Products", "Des données prêtes à l’emploi, documentées, gouvernées et réutilisables par domaine.", ORANGE, "◆")
    card(slide, Inches(4.8), Inches(2.0), Inches(3.75), Inches(3.45), "Produit sémantique certifié", "Une définition commune des KPI, métriques et règles métier, avec qualité et sécurité intégrées.", TEAL, "✓")
    card(slide, Inches(8.9), Inches(2.0), Inches(3.75), Inches(3.45), "Expériences décisionnelles", "Self-service, insights contextualisés et automatisation pour agir plus vite.", VIOLET, "↗")
    text(slide, "UNE CHAÎNE DE VALEUR UNIQUE  →  Données  →  Sens  →  Décisions", Inches(1.0), Inches(6.0), Inches(11.3), Inches(.4), 15, True, WHITE, PP_ALIGN.CENTER)

    slide = base(prs, "Que faut-il maîtriser ?", "03", "Cinq capacités complémentaires pour transformer la donnée en décisions.")
    capabilities = [("Semantic Engineering", "KPI, métriques, règles métier et modèles réutilisables.", TEAL, "01"), ("Delivery & Automatisation", "CI/CD, tests, quality gates, promotion et rollback.", BLUE, "02"), ("Gouvernance & Sécurité", "Standards, certification, lineage, accès et conformité.", VIOLET, "03"), ("Intégration Data", "Préparation, qualité, disponibilité et optimisation des sources.", ORANGE, "04"), ("Analytics avancée", "Forecasting, simulation, ML, Copilots et Data Agents.", GREEN, "05")]
    for i, (h, b, c, n) in enumerate(capabilities):
        x = Inches(.7 + (i % 3) * 4.15)
        y = Inches(1.85 + (i // 3) * 2.25)
        card(slide, x, y, Inches(3.75), Inches(1.8), h, b, c, n)

    slide = base(prs, "Feuille de route en 4 étapes", "04", "Chaque étape délivre de la valeur tout en préparant la suivante.")
    phases = [("2026", "FONDER", "Architecture cible\nStandards de DEV\nPremiers modèles", GREEN), ("2027", "INDUSTRIALISER", "CI/CD et qualité\nSous-domaines\nAdoption", BLUE), ("2028", "ÉTENDRE", "Nouveaux domaines\nSelf-service\nAnalytics avancée", VIOLET), ("2029", "EXCELLER", "IA & Data Agents\nOptimisation\nImpact mesurable", ORANGE)]
    for i, (year, name, body, c) in enumerate(phases):
        x = Inches(.65 + i * 3.12)
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.15), Inches(2.72), Inches(3.35), PANEL, True)
        shape(slide, MSO_SHAPE.OVAL, x + Inches(.92), Inches(1.72), Inches(.85), Inches(.85), c)
        text(slide, year, x, Inches(1.92), Inches(2.72), Inches(.32), 11, True, WHITE, PP_ALIGN.CENTER)
        text(slide, name, x + Inches(.18), Inches(2.85), Inches(2.36), Inches(.4), 14, True, c, PP_ALIGN.CENTER)
        text(slide, body, x + Inches(.28), Inches(3.55), Inches(2.16), Inches(1.25), 11, False, MUTED, PP_ALIGN.CENTER)
        if i < 3:
            text(slide, "→", x + Inches(2.75), Inches(3.55), Inches(.35), Inches(.3), 20, True, TEAL, PP_ALIGN.CENTER)

    slide = base(prs, "Principes de réussite", "05", "Les convictions qui guident nos décisions et notre exécution.")
    principles = [("Centré produit sémantique", "Un socle commun de valeur durable.", TEAL), ("Qualité & Gouvernance", "La confiance, prérequis de l’adoption.", BLUE), ("Adoption & Accompagnement", "Des utilisateurs engagés dans le changement.", VIOLET), ("Agilité & Itération", "Livrer, apprendre et s’améliorer en continu.", ORANGE), ("Collaboration & Partenariat", "IT, Data et métiers unis vers l’impact.", GREEN)]
    for i, (h, b, c) in enumerate(principles):
        card(slide, Inches(.7 + (i % 3) * 4.15), Inches(1.85 + (i // 3) * 2.25), Inches(3.75), Inches(1.8), h, b, c, "★")

    slide = base(prs, "Bénéfices pour l’entreprise", "06", "Une BI industrialisée transforme la donnée en levier de performance.")
    benefits = [("Décisions plus rapides", "Réduction du time-to-insight", TEAL), ("Qualité & Confiance", "Données fiables, sécurisées, traçables", BLUE), ("Efficacité opérationnelle", "Automatisation et réutilisation", ORANGE), ("Expérience utilisateur", "Autonomie et self-service", VIOLET), ("Innovation & scalabilité", "Nouvelles capacités à l’échelle", GREEN), ("Impact business mesurable", "Valeur et ROI objectivés", TEAL)]
    for i, (h, b, c) in enumerate(benefits):
        card(slide, Inches(.7 + (i % 3) * 4.15), Inches(1.8 + (i // 3) * 2.15), Inches(3.75), Inches(1.7), h, b, c, "✓")

    slide = base(prs, "Leviers et organisation", "07", "Un cadre clair pour accélérer, aligner et faire durer la transformation.")
    card(slide, Inches(.7), Inches(1.85), Inches(5.85), Inches(3.9), "Leviers clés", bullets(["Sponsoring engagé", "Vision partagée", "Données maîtrisées", "Compétences développées", "Méthodologie commune"]), TEAL, "↗")
    card(slide, Inches(6.8), Inches(1.85), Inches(5.85), Inches(3.9), "Rôles complémentaires", bullets(["Comité de pilotage", "Product Owner", "Équipe Data & BI", "Data Steward", "Utilisateurs métiers"]), VIOLET, "◎")

    slide = base(prs, "Indicateurs de succès", "08", "Mesurer l’adoption, la qualité et la valeur pour piloter la trajectoire.")
    metrics = [("Adoption", "% utilisateurs actifs mensuels", TEAL), ("Qualité", "% données conformes", BLUE), ("Performance", "Réduction du time-to-insight", ORANGE), ("Fiabilité", "Disponibilité des plateformes", VIOLET), ("Valeur métier", "Cas d’usage à impact", GREEN), ("ROI", "Gains réalisés / investissements", TEAL)]
    for i, (h, b, c) in enumerate(metrics):
        card(slide, Inches(.7 + (i % 3) * 4.15), Inches(1.8 + (i // 3) * 2.15), Inches(3.75), Inches(1.7), h, b, c, "●")

    slide = base(prs, "Prochaines étapes", "09", "Transformer l’ambition en initiatives concrètes, à fort impact.")
    steps = [("1", "Atelier de lancement", "Valider vision et priorités"), ("2", "Cadrage & Gouvernance", "Définir rôles et règles"), ("3", "Plan détaillé", "Construire backlog et trajectoire"), ("4", "Kick-off des initiatives", "Démarrer les premiers chantiers")]
    for i, (n, h, b) in enumerate(steps):
        y = Inches(1.75 + i * 1.15)
        shape(slide, MSO_SHAPE.OVAL, Inches(1.0), y, Inches(.65), Inches(.65), [TEAL, BLUE, VIOLET, ORANGE][i])
        text(slide, n, Inches(1.0), y + Inches(.08), Inches(.65), Inches(.3), 15, True, WHITE, PP_ALIGN.CENTER)
        text(slide, h, Inches(2.0), y, Inches(3.8), Inches(.3), 17, True)
        text(slide, b, Inches(6.1), y, Inches(5.4), Inches(.3), 13, False, MUTED)
        if i < 3:
            shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.29), y + Inches(.65), Inches(.07), Inches(.5), MUTED)

    slide = base(prs, "Passons à l’action", "10", "Une ambition partagée, une valeur durable.")
    text(slide, "La donnée est notre levier.\nLa performance est notre destination.", Inches(1.3), Inches(1.65), Inches(10.7), Inches(1.25), 30, True, WHITE, PP_ALIGN.CENTER)
    actions = ["Adoptons une vision commune et ambitieuse", "Priorisons les initiatives à fort impact", "Collaborons et partageons nos connaissances", "Mesurons nos progrès et célébrons les succès", "Restons agiles et tournés vers l’avenir"]
    for i, item in enumerate(actions):
        text(slide, "✓", Inches(2.0), Inches(3.35 + i * .48), Inches(.35), Inches(.25), 13, True, TEAL)
        text(slide, item, Inches(2.42), Inches(3.35 + i * .48), Inches(8.8), Inches(.25), 13, False, MUTED)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(6.05), Inches(4.35), Inches(.58), ORANGE, True)
    text(slide, "ENSEMBLE, CRÉONS DE LA VALEUR", Inches(4.65), Inches(6.16), Inches(4.05), Inches(.25), 11, True, WHITE, PP_ALIGN.CENTER)

    prs.save(OUT)


if __name__ == "__main__":
    make()
